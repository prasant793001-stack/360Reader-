"""Creates small Office files used by the on-device test."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "360Reader test deck"
s.placeholders[1].text = "If you can read this, PowerPoint rendering works."
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Shapes and colours"
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(5), Inches(3))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x4F, 0x46, 0xE5)
box.text_frame.text = "Calibri text in a shape"
for p in box.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(28); r.font.name = "Calibri"; r.font.color.rgb = RGBColor(255, 255, 255)
prs.save("sample.pptx")

wb = Workbook(); ws = wb.active; ws.title = "Sales"
ws.append(["Month", "Revenue", "Cost"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F9D55")
for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
    ws.append([m, 1000 + i * 150, 600 + i * 80])
wb.save("sample.xlsx")
